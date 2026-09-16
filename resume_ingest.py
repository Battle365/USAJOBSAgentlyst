from __future__ import annotations

import hashlib
import io
import re
import socket
import struct
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import PurePath

from models import EvidenceUnit, ResumeRecord


MAX_RESUME_BYTES = 8 * 1024 * 1024
MAX_PASTED_RESUME_CHARS = 100_000
MIN_EXTRACTED_CHARS = 80
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
EICAR_MARKER = b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!"


class ResumeValidationError(ValueError):
    pass


class MalwareScannerUnavailable(RuntimeError):
    pass


def scan_with_clamav(data: bytes, host: str, port: int = 3310, timeout: float = 10.0) -> bool:
    """Return True only when a ClamAV clamd INSTREAM scan explicitly reports OK."""
    try:
        with socket.create_connection((host, port), timeout=timeout) as connection:
            connection.sendall(b"zINSTREAM\0")
            for offset in range(0, len(data), 64 * 1024):
                chunk = data[offset:offset + 64 * 1024]
                connection.sendall(struct.pack("!I", len(chunk)) + chunk)
            connection.sendall(struct.pack("!I", 0))
            response = connection.recv(4096).decode("utf-8", errors="replace")
    except (OSError, TimeoutError) as exc:
        raise MalwareScannerUnavailable("The upload safety service is unavailable.") from exc
    if " FOUND" in response:
        return False
    if response.rstrip("\0\r\n").endswith(" OK"):
        return True
    raise MalwareScannerUnavailable("The upload safety service returned an invalid response.")


def _safe_filename(filename: str) -> str:
    name = PurePath(filename or "resume").name
    return re.sub(r"[^A-Za-z0-9._ -]", "_", name)[:120]


def _validate_common(data: bytes, filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ResumeValidationError("Upload a PDF or DOCX resume.")
    if not data:
        raise ResumeValidationError("The uploaded file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeValidationError(f"The resume exceeds the {MAX_RESUME_BYTES // 1024 // 1024} MB limit.")
    if EICAR_MARKER in data or b"<script" in data.lower():
        raise ResumeValidationError("The upload was rejected by the file safety check.")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ResumeValidationError("The file content does not match the PDF extension.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise ResumeValidationError("The file content does not match the DOCX extension.")
    return suffix


def _extract_pdf(data: bytes) -> list[tuple[int, str]]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeValidationError("Password-protected or encrypted PDFs are not supported.")
        return [(index, page.extract_text() or "") for index, page in enumerate(reader.pages, start=1)]
    except ResumeValidationError:
        raise
    except Exception as exc:
        raise ResumeValidationError("The PDF could not be parsed. Upload a readable replacement.") from exc


def _extract_docx(data: bytes) -> list[tuple[int, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            names = {item.filename.lower() for item in infos}
            if "word/document.xml" not in names:
                raise ResumeValidationError("The DOCX is malformed or missing its document content.")
            if any("vbaproject" in name or name.endswith((".exe", ".js", ".vbs")) for name in names):
                raise ResumeValidationError("DOCX files containing active or executable content are not supported.")
            total = sum(item.file_size for item in infos)
            compressed = max(1, sum(item.compress_size for item in infos))
            if total > 40 * 1024 * 1024 or total / compressed > 100:
                raise ResumeValidationError("The DOCX archive failed the file safety check.")
        from docx import Document

        document = Document(io.BytesIO(data))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return [(1, "\n".join(blocks))]
    except ResumeValidationError:
        raise
    except Exception as exc:
        raise ResumeValidationError("The DOCX could not be parsed. Upload a readable replacement.") from exc


SECTION_NAMES = {
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "education": "education",
    "skills": "skills",
    "technical skills": "skills",
    "certifications": "certifications",
    "licenses": "licenses",
    "training": "training",
    "clearance": "clearances",
}


def _months(value: str, now: datetime) -> int | None:
    pattern = re.compile(
        r"(?P<sm>Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?|0?[1-9]|1[0-2])?[ /-]*(?P<sy>(?:19|20)\d{2})\s*(?:-|–|—|to)\s*(?:(?P<em>Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?|0?[1-9]|1[0-2])?[ /-]*(?P<ey>(?:19|20)\d{2})|(?P<present>present|current))",
        re.I,
    )
    match = pattern.search(value)
    if not match:
        return None
    month_names = {name.lower(): index for index, name in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}
    def month(raw: str | None) -> int:
        if not raw:
            return 1
        return int(raw) if raw.isdigit() else month_names[raw[:3].lower()]
    start = int(match["sy"]) * 12 + month(match["sm"])
    end = now.year * 12 + now.month if match["present"] else int(match["ey"]) * 12 + month(match["em"])
    return max(0, end - start + 1)


def _evidence_units(pages: list[tuple[int, str]]) -> tuple[EvidenceUnit, ...]:
    units: list[EvidenceUnit] = []
    section = "unspecified"
    active_duration: int | None = None
    for page, text in pages:
        chunks = [re.sub(r"\s+", " ", chunk).strip() for chunk in re.split(r"\n+|(?<=[.!?])\s+(?=[A-Z])", text)]
        for chunk in chunks:
            if not chunk:
                continue
            heading = re.sub(r"[^a-z ]", "", chunk.lower()).strip()
            if heading in SECTION_NAMES and len(chunk) < 40:
                section = SECTION_NAMES[heading]
                active_duration = None
                continue
            stated_duration = _months(chunk, datetime.now(UTC))
            if stated_duration is not None:
                active_duration = stated_duration
            # A dated position heading commonly precedes its bullet duties. Carrying
            # that duration within the same resume section links work to its role;
            # a new dated heading replaces it rather than summing overlapping jobs.
            units.append(EvidenceUnit(chunk[:1000], section, page, duration_months=stated_duration if stated_duration is not None else active_duration))
    return tuple(units)


def parse_resume(data: bytes, filename: str, mime_type: str, owner_id: str, *, malware_scanner=None, require_malware_scan: bool = False) -> ResumeRecord:
    suffix = _validate_common(data, filename)
    if malware_scanner is not None:
        try:
            if not malware_scanner(data):
                raise ResumeValidationError("The upload was rejected by the malware scanner.")
        except ResumeValidationError:
            raise
        except Exception as exc:
            raise ResumeValidationError("The upload safety service is unavailable. Please try again later.") from exc
    elif require_malware_scan:
        raise ResumeValidationError("Resume uploads are temporarily unavailable because the required malware scanner is not configured.")
    pages = _extract_pdf(data) if suffix == ".pdf" else _extract_docx(data)
    text = "\n\n".join(page_text.strip() for _, page_text in pages if page_text.strip()).strip()
    if len(re.sub(r"\s+", "", text)) < MIN_EXTRACTED_CHARS:
        raise ResumeValidationError("Too little readable text was extracted. Upload a text-based PDF or DOCX.")
    if len(re.findall(r"[A-Za-z]{2,}", text)) < 15:
        raise ResumeValidationError("The resume appears image-only or unreadable. Upload a text-based replacement.")
    digest = hashlib.sha256(data).hexdigest()
    return ResumeRecord(
        file_id=str(uuid.uuid4()), owner_id=owner_id, filename=_safe_filename(filename),
        mime_type=mime_type or ("application/pdf" if suffix == ".pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        size=len(data), uploaded_at=datetime.now(UTC).isoformat(), content_hash=digest,
        extracted_text=text, evidence=_evidence_units(pages),
    )


def parse_pasted_resume(value: str, owner_id: str) -> ResumeRecord:
    text = value.strip()
    if len(text) > MAX_PASTED_RESUME_CHARS:
        raise ResumeValidationError("The pasted résumé exceeds the supported size limit.")
    if len(re.sub(r"\s+", "", text)) < MIN_EXTRACTED_CHARS or len(re.findall(r"[A-Za-z]{2,}", text)) < 15:
        raise ResumeValidationError("Paste a readable résumé with work and education details, not just a title or skill list.")
    data = text.encode("utf-8")
    return ResumeRecord(
        file_id=str(uuid.uuid4()), owner_id=owner_id, filename="Pasted résumé", mime_type="text/plain",
        size=len(data), uploaded_at=datetime.now(UTC).isoformat(), content_hash=hashlib.sha256(data).hexdigest(),
        extracted_text=text, evidence=_evidence_units([(1, text)]),
    )
